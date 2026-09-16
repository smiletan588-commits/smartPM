"""Generate week-four reports using the existing week-three Word styles."""
from pathlib import Path
from shutil import copy2
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
ARCHIVE = Path(r'D:\Project Management\Project_Management')
PERIOD = '2026年9月7日 — 2026年9月11日'


def document(template, title, prefix):
    doc = Document(ARCHIVE / template)
    body = doc._element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'):
            body.remove(child)
    doc.core_properties.title = title
    doc.core_properties.subject = 'SmartPM 第四周实践记录'
    doc.core_properties.author = '谭英棋'
    doc.add_heading(title, 0).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(prefix + PERIOD).alignment = WD_ALIGN_PARAGRAPH.CENTER
    if prefix == '周期：':
        doc.add_paragraph('项目：智能项目管理与看板系统（SmartPM）').alignment = WD_ALIGN_PARAGRAPH.CENTER
    return doc


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, 'List Bullet')


def table(doc, headers, rows, ratios):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.autofit = False
    section = doc.sections[0]
    available = (section.page_width - section.left_margin - section.right_margin) / 360000
    widths = [Cm(available * x / sum(ratios)) for x in ratios]
    for col, width in zip(t.columns, widths):
        col.width = width
    for row_data in rows:
        t.add_row()
    for index, values in enumerate([headers] + rows):
        row = t.rows[index]
        props = row._tr.get_or_add_trPr()
        props.append(OxmlElement('w:cantSplit'))
        if index == 0:
            props.append(OxmlElement('w:tblHeader'))
        for cell, value, width in zip(row.cells, values, widths):
            cell.width = width
            cell.text = str(value)
            if index == 0:
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), '6366F1')
                shading.set(qn('w:val'), 'clear')
                cell._tc.get_or_add_tcPr().append(shading)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(4)
                if index == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(10 if index == 0 else 9)
                    if index == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph('')


weekly = document('第三周周报.docx', 'SmartPM 智能协同看板系统 — 第四周周报', '周期：')
weekly.add_heading('一、本周工作概述', 1)
weekly.add_paragraph('本周在第三周成员权限与启动方案的基础上，完成 Docker 环境故障修复、八个页面的响应式视觉重构，并围绕毕业设计补齐任务评论、活动动态、通知中心、可解释风险和 AI 使用评估。进一步实现 CPM 关键路径、不可变计划基线和 What-if 延期模拟，形成从任务协作到排期分析的完整功能链路。')
weekly.add_paragraph('记录依据：project_history.md 中 9 月 7 日、9 月 9 日和 9 月 10 日的开发与验收记录。9 月 8 日、9 月 11 日未发现独立开发记录，保留日期栏目并注明；本周测试结果引用已有验收记录。')

weekly.add_heading('2026年9月7日（周一）— Docker 故障修复与全站视觉重构', 1)
weekly.add_heading('1. Docker 启动环境与脚本修复', 2)
table(weekly, ['序号', '任务', '产出/修改位置', '说明'], [
    ['1', '迁移 Docker 数据盘', 'Docker WSL 数据目录', '将数据盘迁至 D 盘，通过目录联接保持原路径兼容；重启后 Docker 恢复运行。'],
    ['2', '处理数据库端口冲突', 'docker-compose.yml', '宿主机端口调整为 3307，容器内仍使用 3306，避开本机 MySQL80。'],
    ['3', '修复虚拟化误判', 'start.bat', '综合 HypervisorPresent 与固件虚拟化标志判断，避免 Hyper-V/WSL2 已运行时误报。'],
], [1, 2.4, 3.2, 6])
weekly.add_heading('2. 全站设计系统与响应式适配', 2)
bullets(weekly, [
    '建立冷灰浅色背景与钴蓝品牌色，统一品牌标志、侧栏、页面标题、空状态和加载骨架，覆盖登录、项目列表、任务看板、数据分析、Wiki、项目管理、系统管理和回收站八个路由。',
    '手机看板采用状态分段切换；Wiki 增加文档抽屉；甘特图限制内部滚动；管理员与回收站页面在手机端采用记录卡片。',
    '前端构建通过；Playwright 在 1440、1024、768、390 四档宽度完成 32 组页面检查，无根节点水平溢出或运行时错误；新版前端已部署至 Docker Nginx。',
])

weekly.add_heading('2026年9月8日（周二）— 开发记录说明', 1)
weekly.add_paragraph('现有项目历史中未发现当天单独记录，不将其他日期的实现或测试归入当天。周一已交付的部署与界面改造成果统一列入本周汇总。')

weekly.add_heading('2026年9月9日（周三）— 代码归档与毕业设计升级', 1)
table(weekly, ['序号', '任务', '实现内容'], [
    ['1', 'GitHub 版本归档', '检查版本库边界与忽略规则，排除密钥、依赖和构建产物；稳定版本提交并推送至 main，建立远程跟踪关系。'],
    ['2', '综合升级方案', '明确评论、动态、通知、可解释风险、AI 使用评估、数据库迁移与自动化测试的实施范围。'],
    ['3', '后端数据与接口骨架', '接入 Flyway、环境化 JWT/CORS、真实完成时间、关系型任务依赖和里程碑关联；补齐协作、通知、风险与 AI 日志模块。'],
], [1, 3, 8.6])
weekly.add_paragraph('阶段验证：后端执行 mvn -q -DskipTests compile 通过，为后续前后端联调建立基础。')

weekly.add_heading('2026年9月10日（周四）— 协作、风险、AI 评估与排期升级', 1)
weekly.add_heading('1. 协作闭环与 AI 效果评估', 2)
table(weekly, ['模块', '关键实现', '前端入口/产出'], [
    ['评论与动态', '任务评论、显式成员提醒、只读活动时间线；WebSocket 推送协作变化。', '任务详情评论与动态区域'],
    ['通知中心', '全局通知抽屉，每小时扫描到期、逾期和阻塞情况，支持通知去重。', '应用侧栏壳层通知入口'],
    ['风险中心', '六条确定性规则评分、成员工作量与依赖阻塞分析，结合 Redis 缓存及 AI 结构化建议。', '项目管理页风险中心'],
    ['AI 使用评估', '记录场景、模型、耗时、生成量、采纳情况及 1—5 分反馈，展示五项 AI 指标。', '分析页及 AI 操作日志'],
    ['实验材料', '准备 20 组项目需求、评价指南和结果 CSV 模板；尚未形成真实实验结果。', 'docs 下评估数据集、指南与模板'],
], [2.2, 7, 3.4])
weekly.add_heading('2. CPM 关键路径、计划基线与 What-if 模拟', 2)
bullets(weekly, [
    '新增纯内存 ScheduleEngine，采用拓扑排序及 CPM 前向/后向计算，输出最早/最晚日期、总浮动、关键任务及代表性关键路径，复杂度为 O(V+E)。',
    '仅计算未删除主任务，使用完成—开始依赖和日历日。工期优先取起止日期含首尾天数；日期缺失时按预计工时向上折算，无工时则按 1 天推算；显式提示循环依赖、范围外依赖和日期冲突。',
    'Flyway V4 新增命名基线与不可变任务快照，支持新增、移除、更新和未变化四类偏差；历史任务软删除后仍保留快照。',
    '实现单任务延期模拟，展示项目完工日、受影响任务和风险分布变化；模拟不回写真实任务、依赖或基线。VIEWER 可查看和模拟，MEMBER 可创建基线，项目负责人和 PROJECT_ADMIN 可删除基线。',
    '风险规则纳入零浮动关键任务及相对基线延期，缓存随任务或基线变化失效；甘特图用灰色基线条与红色关键路径展示差异，并增加响应式模拟抽屉。',
])
weekly.add_heading('3. 工程质量与交付验证', 2)
bullets(weekly, [
    '完成 Flyway V1—V4 迁移、OpenAPI 文档、JWT/CORS 环境配置、停用账号 Token 即时失效、密码规则与参数校验，完善容器非 root 运行及构建缓存。',
    '最终阶段后端 14 项测试、前端 9 项测试通过；前端生产构建通过。覆盖空库迁移、权限隔离、依赖阻塞、通知去重、风险评分、基线历史快照和模拟无副作用等场景。',
    '完成四档宽度、八个路由及基线甘特图、What-if 抽屉的浏览器巡检；此前综合升级验收确认四个 Compose 服务健康，首页、健康检查与 Swagger 返回 200。',
])

weekly.add_heading('2026年9月11日（周五）— 周末节点与待办说明', 1)
weekly.add_paragraph('现有项目历史中未发现当天单独开发或测试记录。本周成果以截至 9 月 10 日的已记录验收为准；本周报及 AI 使用记录于 9 月 13 日补充整理。')

weekly.add_heading('二、本周验证与质量情况', 1)
table(weekly, ['类别', '验证结果', '说明'], [
    ['后端自动化测试', '14 项通过', '最终调度升级：8 项调度引擎、4 项风险规则、2 项 MySQL 集成测试。'],
    ['前端测试与构建', '9 项测试通过；构建通过', '最终阶段 npm test -- --run 与 npm run build 均通过。'],
    ['数据库迁移', 'V1—V4 空库迁移成功', '已有库升级至 V3 的部署检查与 V4 空库集成验证分别记录，不混同。'],
    ['浏览器巡检', '四档宽度、八个路由通过', '含新增基线与模拟交互，无根节点横向溢出或运行时错误。'],
    ['容器与接口', '综合升级验收通过', 'MySQL、Redis、后端、Nginx healthy；首页、健康检查、Swagger 返回 200。'],
    ['AI 量化实验', '材料已准备，结果待采集', '20 组需求与结果模板不等同于 20 次已完成的真实模型实验。'],
], [2.5, 3.3, 6.8])
weekly.add_heading('三、本周成果汇总', 1)
table(weekly, ['模块', '关键成果', '状态'], [
    ['部署与界面', '环境冲突修复；八路由统一设计与响应式适配', '已完成'],
    ['协作与通知', '评论、动态、成员提醒、通知扫描和实时推送', '已完成'],
    ['风险与 AI 评估', '确定性风险规则、AI 建议、操作日志和效果指标', '功能完成；实验待开展'],
    ['排期管理', 'CPM、命名基线、偏差比较、无副作用延期模拟', '已完成'],
    ['工程质量', '数据库迁移、接口文档、安全配置和自动化测试', '已完成本周验收'],
], [2.8, 6.5, 3.3])
weekly.add_heading('四、下周计划', 1)
bullets(weekly, [
    '实际运行 20 组 AI 评估需求，采集成功率、耗时、采纳和评分等数据，形成可复现的实验结果与分析。',
    '继续完善 Wiki 多人编辑的实时更新、在线状态与冲突处理方案，明确与任务评论协作的功能边界。',
    '补充关键路径、计划基线和风险规则的毕业设计说明、操作截图及答辩演示案例。',
    '继续开展真实多人使用与较大规模任务依赖场景验证，并核对部署环境数据库迁移版本。',
])

ai = document('第三周ai使用记录.docx', 'SmartPM 项目 — 第四周 AI 使用记录', '记录周期：')
ai.add_heading('一、AI 使用概述', 1)
ai.add_paragraph('本周 AI 辅助工作集中于环境排错、全站 UI 重构、协作功能实现、可解释风险设计、AI 使用评估，以及 CPM 调度与计划基线。开发辅助记录依据项目历史归纳；本次文档由 Codex 辅助整理。历史未逐项记录所用模型、交互次数、Token 与费用，因此不沿用前周的估算轮次，也不追溯指定具体模型。')
ai.add_paragraph('系统内 AI 功能与开发辅助分别记录。AI 负责提供实现建议、代码和说明，最终采纳依据为需求约束、代码检查及已有测试验收；CPM 排期和风险规则使用确定性算法，不由大模型直接决定。')

records = [
    ('二、2026年9月7日 — 环境诊断与响应式界面改造', [
        '问题输入：Docker 启动失败、数据盘占用 C 盘、3306 端口冲突，以及虚拟化已启用仍被脚本误报。',
        'AI 辅助处理：结合 Docker 状态、数据路径、端口占用和虚拟化标志梳理原因，辅助完成数据盘迁移、3307 端口映射及脚本判断修复。',
        '界面辅助：梳理八个路由的布局问题，建立冷灰/钴蓝设计系统和可复用侧栏、标题、状态组件，完善手机看板、Wiki 抽屉与记录卡片。',
        '采纳与验证：以 Docker 恢复运行、构建通过及四档宽度共 32 组页面检查为依据，确认环境修复和界面改造结果。',
    ], '未记录逐轮日志；主要涉及环境诊断、脚本修复、UI 实现与浏览器验证。'),
    ('三、2026年9月8日 — 记录完整性说明', [
        '项目历史未发现当天独立 AI 使用记录，无法确认具体请求、生成内容或交互次数。',
        '不将其他日期的代码生成、调试或验证活动转记为当天成果。',
    ], '无独立记录，不估算。'),
    ('四、2026年9月9日 — 版本归档与综合升级方案', [
        '版本管理辅助：检查项目文件边界与忽略规则，整理稳定版本提交；确保本地密钥和构建产物不进入版本库。',
        '方案输入：在既有权限、看板和 AI 功能基础上提升毕业设计的协作完整性、风险可解释性及实验可验证性。',
        'AI 辅助产出：拆分评论、活动、通知、风险、AI 日志和自动化测试任务，补充数据实体、服务、接口与 Flyway 迁移骨架。',
        '采纳与验证：沿用现有 API 和设计系统，后端编译通过后继续进行功能联调；不以生成代码本身作为完成依据。',
    ], '未记录逐轮日志；主要涉及代码审查、方案设计、后端实现与编译检查。'),
    ('五、2026年9月10日 — 协作实现、AI 评估与排期算法', [
        '协作实现辅助：完成评论、成员提醒、活动时间线、通知扫描及 WebSocket 推送，补齐前端入口和权限限制。',
        '风险与 AI 评估辅助：实现确定性规则评分与结构化建议，记录模型、耗时、生成量、采纳及反馈；编制 20 组项目需求、评价指南和 CSV 模板。',
        '调度算法辅助：依据主任务依赖设计拓扑排序和 CPM 前后向计算，明确日历日、工期推算、循环依赖与日期冲突的处理规则。',
        '基线与模拟辅助：实现不可变命名基线、历史快照、偏差比较和单任务延期模拟；联动风险分布、关键路径甘特图与响应式抽屉。',
        '问题修正：处理未分配任务在基线风险计算中的空键边界，完善缓存失效及默认/指定基线行为。',
        '验证约束：重点检查项目隔离、VIEWER/MEMBER/PROJECT_ADMIN 权限、历史快照保留和模拟不回写；最终阶段后端 14 项、前端 9 项测试通过，构建与浏览器巡检通过。',
    ], '未记录逐轮日志；主要涉及全栈开发、算法设计、边界修复和自动化验证。'),
    ('六、2026年9月11日 — 周度记录说明', [
        '项目历史未发现当天独立 AI 使用记录，本周结果以截至 9 月 10 日的实现与验收记录为准。',
        '第四周两份文档于 9 月 13 日参照前三周格式补充整理，未将此次文档生成计为 9 月 11 日活动。',
    ], '无独立记录，不估算。'),
]
for heading, items, stats in records:
    ai.add_heading(heading, 1)
    ai.add_paragraph('AI 开发辅助任务与核验：')
    bullets(ai, items)
    ai.add_paragraph('AI 交互统计：' + stats)

ai.add_heading('七、系统内 AI 功能使用与变更', 1)
table(ai, ['AI 功能/模块', '本周使用或改进', '结果'], [
    ['风险分析建议', '在规则评分、工作量和依赖分析基础上，提供 AI 结构化建议；结合缓存降低重复计算。', '已实现建议功能'],
    ['AI 操作日志', '统一记录场景、模型、耗时、生成量、采纳情况及 1—5 分反馈。', '形成评估数据采集能力'],
    ['AI 效果展示', '分析页展示五项 AI 指标，与操作日志相衔接。', '页面与统计能力已完成'],
    ['评估实验材料', '20 组不同类型项目需求、评价指南与结果 CSV 模板。', '待实际调用与人工评分'],
    ['既有 AI 业务', '任务拆解、项目总结、Wiki 写作及项目计划等能力作为既有基础保留。', '无本周逐次调用数据，不填调用量'],
    ['CPM 与 What-if', '采用确定性算法计算关键路径、基线偏差及延期影响，AI 辅助其研发。', '属于调度功能，不计为模型调用'],
], [2.7, 6.5, 3.4])
ai.add_heading('八、AI 使用汇总与反思', 1)
table(ai, ['维度', '本周情况', '说明'], [
    ['开发辅助工具与模型', '历史未逐项注明；本次整理使用 Codex', '无法据此前周工具名称推定本周全部研发工具。'],
    ['总交互轮次', '未统计', '缺少完整会话日志，不给出估算总数或累计数。'],
    ['Token 与费用', '未统计', '当前材料不足以还原真实模型消耗与费用。'],
    ['主要产出', '部署、UI、协作、风险、AI 评估及 CPM', '以项目历史中的实现与验收里程碑归纳。'],
    ['质量依据', '后端 14 项、前端 9 项测试及构建通过', '采用本周最终阶段结果，不将中间验收次数重复累加。'],
    ['实验数据', '20 组需求已准备，真实结果待采集', '不宣称已获得准确率、平均响应时间或效率提升比例。'],
], [2.5, 4.5, 5.6])
ai.add_paragraph('使用成效：AI 辅助将跨环境、前后端和数据库的改造拆解为可验证步骤，并支持把排期规则转化为可复现的实现与测试。相较单纯生成代码，本周更注重权限、历史数据、异常边界和无副作用约束。')
ai.add_paragraph('经验与不足：环境问题必须依据实际端口、进程和日志定位；生成的算法实现必须由明确规则和边界用例验证。当前缺少开发会话和模型成本的完整统计，产品内评估也仍需真实实验，暂不能给出量化提效结论。')
ai.add_heading('九、后续 AI 使用计划', 1)
bullets(ai, [
    '按评价指南实际运行 20 组需求，记录输入、模型、时间、生成结果、采纳情况及人工评分，保留失败样本。',
    '完善开发辅助使用台账，逐次记录工具、模型、目的、关键输入、采纳修改及验证结果；可获取时补充 Token 与费用。',
    '整理 Prompt 与评估条件，逐步进行版本管理和对比实验，保证实验可复现。',
    '继续利用 AI 辅助补充多人协作与复杂依赖图用例，排期结论仍由确定性算法及测试确认。',
])

for doc, name in [(weekly, '第四周周报.docx'), (ai, '第四周ai使用记录.docx')]:
    path = ROOT / name
    if path.exists() or (ARCHIVE / name).exists():
        raise FileExistsError(f'Refusing to overwrite existing report: {name}')
    doc.save(path)
    copy2(path, ARCHIVE / name)
    print(path)
